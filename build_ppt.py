"""Build CodeCarto R1-07 Round-1 deck (7 + 1 slides) — editorial paper/ink theme."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

PAPER = RGBColor(0xF2, 0xEB, 0xD8)
INK = RGBColor(0x0E, 0x0E, 0x10)
ASH = RGBColor(0x5C, 0x59, 0x52)
OXBLOOD = RGBColor(0x7A, 0x1C, 0x1C)
CARD = RGBColor(0xF7, 0xF3, 0xE8)
RULE = RGBColor(0xC9, 0xC2, 0xAE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DISPLAY, BODY, MONO = "Georgia", "Calibri", "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color=PAPER):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()
    return r


def bar(slide, y, h, color, x=0, w=None):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, prs.slide_width if w is None else w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()
    return r


def text(slide, x, y, w, h, runs, size=18, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT, space_after=Pt(4)):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, chunk in enumerate(runs):
        if isinstance(chunk, str):
            chunk = {"t": chunk}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = space_after
        r = p.add_run(); r.text = chunk["t"]
        r.font.size = Pt(chunk.get("size", size))
        r.font.bold = chunk.get("bold", bold)
        r.font.color.rgb = chunk.get("color", color)
        r.font.name = chunk.get("font", font)
    return tb


def bullets(slide, x, y, w, h, items, size=17):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after, p.space_before = Pt(6), Pt(2)
        if head:
            r = p.add_run(); r.text = head + "  "
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, OXBLOOD, BODY
        r = p.add_run(); r.text = body
        r.font.size, r.font.color.rgb, r.font.name = Pt(size), INK, BODY
    return tb


def foot(slide, num, label):
    bar(slide, Inches(7.06), Inches(0.06), OXBLOOD)
    text(slide, 0.6, 7.14, 5, 0.3, [{"t": "CodeCarto  ·  R1-07 Secure Software Delivery", "size": 10, "color": ASH, "font": MONO}])
    text(slide, 11.6, 7.14, 1.1, 0.3, [{"t": f"{num} / 8", "size": 10, "color": ASH, "font": MONO}], align=PP_ALIGN.RIGHT)
    text(slide, 0.6, 0.25, 12, 0.3, [{"t": label, "size": 10, "color": ASH, "font": MONO}])


def title(slide, kicker, heading, sub=None):
    text(slide, 0.6, 0.65, 12, 0.5, [{"t": kicker, "size": 13, "bold": True, "color": OXBLOOD}])
    text(slide, 0.6, 1.05, 12, 1.0, [{"t": heading, "size": 40, "bold": True, "font": DISPLAY}])
    if sub:
        text(slide, 0.6, 1.95, 12, 0.6, [{"t": sub, "size": 16, "color": ASH}])


# ---------------- 1 · Introduction ----------------
s = prs.slides.add_slide(BLANK); bg(s)
bar(s, Inches(3.35), Inches(0.08), OXBLOOD)
text(s, 0.6, 0.5, 12, 0.4, [{"t": "HACKROYNX 2.0  ·  ROUND 1  ·  PS ID: R1-07", "size": 14, "bold": True, "color": OXBLOOD, "font": MONO}], align=PP_ALIGN.CENTER)
text(s, 0.6, 1.1, 12, 1.4, [{"t": "WORLD MONITOR", "size": 72, "bold": True, "font": DISPLAY}], align=PP_ALIGN.CENTER)
text(s, 0.6, 2.5, 12, 0.7, [{"t": "A fail-closed security release gate for modern delivery pipelines", "size": 24, "color": ASH, "font": DISPLAY}], align=PP_ALIGN.CENTER)
for i, (n, l) in enumerate([("12", "security scanners"), ("CVSS 3.1", "risk scoring"), ("49", "tests green"), ("4", "report formats")]):
    x = 2.4 + i * 2.35
    text(s, x, 3.9, 2.2, 0.9, [{"t": n, "size": 34, "bold": True, "font": DISPLAY}, {"t": "\n" + l, "size": 13, "color": ASH}], align=PP_ALIGN.CENTER)
    if i < 3:
        d = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 2.28), Inches(3.95), Inches(0.02), Inches(0.8))
        d.fill.solid(); d.fill.fore_color.rgb = RULE; d.line.fill.background()
text(s, 0.6, 5.5, 12, 0.5, [{"t": "Team CodeCarto  ·  scans third-party code, configs, secrets & containers — then BLOCKS unsafe releases", "size": 15, "color": ASH}], align=PP_ALIGN.CENTER)
r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.4), Inches(6.1), Inches(2.5), Inches(0.55))
r.fill.solid(); r.fill.fore_color.rgb = INK; r.line.fill.background()
r.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
run = r.text_frame.paragraphs[0].add_run(); run.text = "▶  Live demo ready"
run.font.size, run.font.color.rgb, run.font.name = Pt(15), WHITE, BODY
foot(s, 1, "01 · INTRODUCTION")

# ---------------- 2 · Problem & Approach ----------------
s = prs.slides.add_slide(BLANK); bg(s)
title(s, "02 · PROBLEM & APPROACH  —  PS ID: R1-07 Secure Software Delivery",
      "Ship fast, but never ship blind.",
      "Pipelines move code, libraries, images, configs and secrets to prod daily — one overlooked flaw becomes breach, outage or supply-chain compromise.")
text(s, 0.6, 2.9, 5.6, 0.4, [{"t": "THE PROBLEM", "size": 13, "bold": True, "color": OXBLOOD, "font": MONO}])
bullets(s, 0.6, 3.3, 5.6, 3.3, [
    ("Vulnerable dependencies", "— known CVEs ride into prod inside third-party libraries."),
    ("Leaked secrets", "— tokens & keys hardcoded across contributors' code."),
    ("Silent misconfiguration", "— weak headers, TLS and cookies nobody grades."),
    ("No release brakes", "— nothing stops a risky build from deploying."),
])
text(s, 7.0, 2.9, 5.7, 0.4, [{"t": "OUR APPROACH — FAIL-CLOSED BY DEFAULT", "size": 13, "bold": True, "color": OXBLOOD, "font": MONO}])
steps = ["SCAN\n12 modules", "NORMALIZE\none schema", "SCORE\nCVSS + health", "GATE\nblock / approve", "RETEST\nprove the fix"]
for i, st in enumerate(steps):
    x = 7.0 + i * 1.12
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.4), Inches(1.0), Inches(1.15))
    b.fill.solid(); b.fill.fore_color.rgb = CARD; b.line.color.rgb = RULE
    tf = b.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = st; r.font.size, r.font.name, r.font.color.rgb = Pt(12), MONO, INK
    if i < 4:
        text(s, x + 1.0, 3.85, 0.15, 0.3, [{"t": "›", "size": 20, "bold": True, "color": OXBLOOD}], align=PP_ALIGN.CENTER)
bullets(s, 7.0, 4.85, 5.7, 1.8, [
    ("", "Every build is scanned, scored 0–100, and gated: BLOCKED until criticals, health floor and evidence checks pass."),
    ("", "Developers get fix guidance plus proof-of-fix — safety without slowing the pipeline."),
])
foot(s, 2, "02 · PROBLEM & APPROACH")

# ---------------- 3 · Innovation & Solution ----------------
s = prs.slides.add_slide(BLANK); bg(s)
title(s, "03 · INNOVATION & SOLUTION", "Not another scanner — a decision engine.")
items = [
    ("01 · One schema for every signal", "SAST, DAST, SCA, secrets and config findings normalize into a single Common Finding Format with CVSS 3.1 — apples-to-apples risk ranking."),
    ("02 · Evidence you can trust", "Tokens and keys masked before storage, filesystem scans jailed, DNS pinned, SSRF blocked — results are reproducible, not hallucinated."),
    ("03 · Policy as a release gate", "Block on criticals, health floor, incomplete scans. CI-native: quality gates, SARIF upload, CycloneDX SBOM out."),
    ("04 · Proof, not promises", "Fingerprint-based retest (sha1 of target + check + component) verdicts each fix FIXED or STILL PRESENT, with before/after health delta."),
]
for i, (h, b) in enumerate(items):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.2, 2.7 + row * 2.15
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.9), Inches(1.9))
    card.fill.solid(); card.fill.fore_color.rgb = CARD; card.line.color.rgb = RULE
    text(s, x + 0.3, y + 0.2, 5.3, 1.5, [{"t": h, "size": 17, "bold": True, "color": OXBLOOD}, {"t": "\n" + b, "size": 14}])
foot(s, 3, "03 · INNOVATION & SOLUTION")

# ---------------- 4 · Technical Architecture ----------------
s = prs.slides.add_slide(BLANK); bg(s)
title(s, "04 · TECHNICAL ARCHITECTURE", "Pipeline in, decision out.")
layers = [
    ("TARGETS", "vulnerable lab :8080  ·  real app :3000  ·  source tree"),
    ("AUTH GATE", "loopback / RFC1918 only  ·  SSRF + metadata-IP blocked  ·  RBAC viewer / analyst / admin"),
    ("ORCHESTRATOR", "thread-per-assessment  ·  watchdog + timeouts  ·  12 scanner modules"),
    ("FINDING ENGINE", "normalize → dedupe (fingerprint) → CVSS 3.1 → mask evidence"),
    ("DECIDE", "policy gate BLOCKED / APPROVED  ·  PDF · JSON · MD · CSV  ·  retest loop"),
]
y = 2.55
for head, body in layers:
    b = bar(s, Inches(y), Inches(0.78), CARD, x=Inches(0.6), w=Inches(12.13))
    b.line.color.rgb = RULE
    text(s, 0.85, y + 0.07, 2.4, 0.65, [{"t": head, "size": 14, "bold": True, "color": OXBLOOD, "font": MONO}])
    text(s, 3.4, y + 0.16, 9.1, 0.5, [{"t": body, "size": 14, "font": MONO}])
    y += 0.88
text(s, 0.6, 7.0 - 0.62, 12.1, 0.4, [{"t": "FastAPI  ·  SQLAlchemy / SQLite  ·  vanilla-JS SPA  ·  Go binaries (portia · bomber · chainscanner)  ·  Docker non-root  ·  JWT + audit logs", "size": 12, "color": ASH, "font": MONO}], align=PP_ALIGN.CENTER)
foot(s, 4, "04 · TECHNICAL ARCHITECTURE")

# ---------------- 5 · Implementation Details ----------------
s = prs.slides.add_slide(BLANK); bg(s)
title(s, "05 · IMPLEMENTATION DETAILS", "The math and machinery behind the verdict.")
bullets(s, 0.6, 2.8, 5.9, 3.9, [
    ("Health score", "100 − (5×Critical + 3×High + 1.5×Medium + 0.5×Low), clamped 0–100. 38 curated CVSS FIRST-vector presets."),
    ("Retest proof", "Re-runs only the failing check; fingerprint compare → FIXED / STILL_PRESENT / INCONCLUSIVE (never false FIXED)."),
    ("Evidence safety", "Secrets/cookies/keys redacted pre-write; scanners jailed to authorized paths; every run audit-logged."),
])
bullets(s, 6.9, 2.8, 5.8, 3.9, [
    ("Policy knobs", "block_on_critical, block_on_high count, health floor, block_on_incomplete — fail-closed defaults."),
    ("Access control", "viewer reads · analyst scans/retests · admin audits. Registration gate via REGISTRATION_ENABLED."),
    ("Scale path", "server-side severity + text search, limit/offset, X-Total-Count — findings UI pages 25 at a time."),
])
foot(s, 5, "05 · IMPLEMENTATION DETAILS")

# ---------------- 6 · Feasibility & Impact ----------------
s = prs.slides.add_slide(BLANK); bg(s)
title(s, "06 · FEASIBILITY & IMPACT", "Runs in a minute. Pays off on every release.")
bullets(s, 0.6, 2.8, 5.9, 3.9, [
    ("60-second start", "pip install → start_all.py → lab :8080 + platform :8000. Dockerized, CI on every push."),
    ("Proven quality", "49 pytest green · pip-audit · SARIF upload · quality + release gates in CI."),
    ("Real coverage", "OWASP Top-10 classes: auth, IDOR, SQLi, XSS, headers, TLS, secrets, CVEs, supply chain."),
])
bullets(s, 6.9, 2.8, 5.8, 3.9, [
    ("Impact", "Criticals blocked pre-deploy · devs get fix + proof · releases ship with evidence, not hope."),
    ("No slowdown", "Only failing checks re-run; everything else is cached verdicts."),
    ("Roadmap", "Postgres · OPA policies · IDE plugin · scheduled pipeline scans."),
])
foot(s, 6, "06 · FEASIBILITY & IMPACT")

# ---------------- 7 · Team ----------------
s = prs.slides.add_slide(BLANK); bg(s)
bar(s, Inches(3.35), Inches(0.08), OXBLOOD)
text(s, 0.6, 0.5, 12, 0.4, [{"t": "07 · TEAM", "size": 14, "bold": True, "color": OXBLOOD, "font": MONO}], align=PP_ALIGN.CENTER)
text(s, 0.6, 0.95, 12, 1.0, [{"t": "CodeCarto", "size": 64, "bold": True, "font": DISPLAY}], align=PP_ALIGN.CENTER)
members = [
    ("Mehul Wagde", "linkedin.com/in/mehul-wagde-750796276"),
    ("Parth Toshniwal", "linkedin.com/in/parth-toshniwal-807211435"),
    ("Sarthak Sute", "linkedin.com/in/sarthak-sute"),
]
for i, (name, link) in enumerate(members):
    x = 1.15 + i * 3.9
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.6), Inches(3.3), Inches(1.9))
    card.fill.solid(); card.fill.fore_color.rgb = CARD; card.line.color.rgb = RULE
    text(s, x + 0.25, 2.85, 2.8, 1.4, [
        {"t": name, "size": 21, "bold": True, "font": DISPLAY},
        {"t": "\n" + link, "size": 11, "color": ASH, "font": MONO},
    ], align=PP_ALIGN.CENTER)
text(s, 0.6, 5.0, 12, 0.9, [
    {"t": "Faculty Mentor  ·  ", "size": 15, "bold": True, "color": OXBLOOD},
    {"t": "Shringarika Pandey", "size": 20, "bold": True, "font": DISPLAY},
], align=PP_ALIGN.CENTER)
text(s, 0.6, 5.85, 12, 0.4, [{"t": "PS ID: R1-07 Secure Software Delivery  ·  HackRoynx 2.0 Round 1", "size": 13, "color": ASH, "font": MONO}], align=PP_ALIGN.CENTER)
foot(s, 7, "07 · TEAM MEMBERS & MENTOR")

# ---------------- 8 · Extra: demo storyboard ----------------
s = prs.slides.add_slide(BLANK); bg(s)
title(s, "APPENDIX · LIVE DEMO STORYBOARD", "Sixty seconds from scan to proven fix.")
steps = [
    ("1 · SCAN", "New Assessment on the lab playground — 12 modules, live progress."),
    ("2 · TRIAGE", "Findings ranked by CVSS; gate reads BLOCKED on criticals."),
    ("3 · FIX", "Flip lab patch toggles (headers, SQLi, IDOR…) — the actual code path heals."),
    ("4 · PROVE", "Retest → FIXED overlay; health climbs 68 → 91 with evidence attached."),
]
for i, (h, b) in enumerate(steps):
    x = 0.6 + i * 3.08
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.7), Inches(2.9), Inches(2.2))
    card.fill.solid(); card.fill.fore_color.rgb = CARD; card.line.color.rgb = RULE
    text(s, x + 0.3, 2.95, 2.3, 1.7, [{"t": h, "size": 16, "bold": True, "color": OXBLOOD, "font": MONO}, {"t": "\n" + b, "size": 14}])
text(s, 0.6, 5.3, 12.1, 0.8, [{"t": "Swap these boxes for real screenshots before submitting — judges score what they can see.", "size": 14, "color": ASH}], align=PP_ALIGN.CENTER)
foot(s, 8, "APPENDIX · DEMO")

out = "CodeCarto_R1-07_Secure_Software_Delivery.pptx"
prs.save(out)
print("saved", out, "| slides:", len(prs.slides._sldIdLst))
